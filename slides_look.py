#!/usr/bin/env python3
"""The deck "look": chrome knobs + the shapes that draw them, in one place.

This is the single source of truth for the visual chrome the
`layout_to_pptx.py` renderer applies. Keeping the knobs and the
oxml-fragment builders here means "tweak the look" is one file.

Chrome here is the deck-level furniture that isn't slide *content*: a faint
boundary frame, an auto-updating slide number, and a barely-there outline on
content text boxes (so editing bounds read). Defaults (slide size, fonts,
margins) for the native renderer live alongside, so a terse layout file inherits
a sensible look without restating every knob.

Built on **python-pptx**: shapes go on the parsed object model. The two elements
the high-level API can't create — a shape on the slide *master* and an
auto-updating slide-number *field* — are built as oxml fragments via `parse_xml`.
"""

from __future__ import annotations

from dataclasses import dataclass

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.slide import Slide, SlideMaster
from pptx.util import Emu, Pt

# --- look knobs (tweak here) -----------------------------------------------
# Slide canvas, in EMU. 16:9 at 10in wide. 914400 EMU = 1 inch.
SLIDE_W, SLIDE_H = 9144000, 5143500

# Faint boundary frame on the master, so the slide edge reads on a white deck.
BORDER_HEX = "C2C8D2"
BORDER_WEIGHT = 6350                  # EMU -> 0.5pt
BORDER_INSET = 36576                  # ~1mm in from each edge so viewers keep it

# Auto-updating slide number, bottom-right.
SLIDE_NUM_HEX = "808080"
SLIDE_NUM_SZ = 1100                   # hundredths of a point -> 11pt
SLIDE_NUM_OFF = (8229600, 4654550)    # EMU; bottom-right corner
SLIDE_NUM_EXT = (685800, 365125)      # EMU; box size

# Title type: left-aligned + bold + down-sized from PowerPoint's big default.
TITLE_SZ = 2000                       # hundredths of a point -> 20pt
# Wrapped titles collide in Google Slides at a flat 100% multiple (spcPct is scaled
# off the font's own metrics, so a big bold title crowds). Default every title to
# EXACT point line spacing at this multiple of its font size (spcPts), which Slides
# honours literally so a 2-line title never overlaps — e.g. 1.25 -> 40pt on a 32pt
# title. Author override: line-spacing= on the <title> (a float multiple, or "40pt").
TITLE_LINE_SPACING_FACTOR = 1.25

# Faint outline on every content text box so editing boundaries are visible.
# Toggle off to present a clean deck with no box outlines at all.
TEXTBOX_OUTLINE_SHOW = False          # draw the faint per-box editing outline?
TEXTBOX_OUTLINE_HEX = "EFF1F4"        # barely-there, just off-white
TEXTBOX_OUTLINE_WEIGHT = 3175         # EMU -> 0.25pt

# Native-renderer body defaults (the legacy path inherits these from the
# layout/master instead). Point sizes; the renderer scales EMU itself.
BODY_SZ_PT = 14                       # default body / bullet size
BODY_COLOR_HEX = "1A1A1A"             # near-black body text
TITLE_COLOR_HEX = "1A1A1A"
# `str`, not the inferred literal — the renderer overrides it from <deck font=…>.
FONT_FAMILY: str = "Calibri"          # PowerPoint's default sans; safe in Slides
# Default content-box geometry (a comfortable title + body frame, in inches),
# used when a layout file leaves a shape's box implicit.
DEFAULT_TITLE_BOX_IN = (0.5, 0.3, 9.0, 0.9)     # x, y, w, h
DEFAULT_BODY_BOX_IN = (0.5, 1.3, 9.0, 3.9)
# Default corner radius (inches) for <shape> rounded-rectangles. Small on purpose:
# python-pptx's ~0.167-of-shorter-side default is a fat curve that clips top-anchored
# text into the corner. radius= on the element overrides this per-shape.
SHAPE_RADIUS_IN = 0.05

# Caption type: a small, muted, italic contextual note — a figure caption or a
# side annotation. One source for the house caption look, folding the recurring
# ad-hoc `<text size="12" color="555555"><i>…</i></text>` into the <caption>
# element so authors don't restate it (and it stays consistent deck-to-deck).
CAPTION_SZ_PT = 12
CAPTION_COLOR_HEX = "555555"          # muted grey (matches the section-header grey)
CAPTION_ITALIC = True

# Provenance link (<source>): a small, muted corner link from a data slide back to
# where its data lives. Set SOURCE_BASE_URL (a report server, a repo URL, ...) and
# `<source run="path"/>` resolves against it; leave it empty and every <source>
# must state a full href=. One place for the scheme + look.
SOURCE_BASE_URL = ""
SOURCE_SUMMARY_FILE = ""  # optional file appended to run= (e.g. "index.html")
SOURCE_LABEL = "data ↗"
SOURCE_SZ_PT = 9
SOURCE_COLOR_HEX = "8A8A8A"           # faint grey; a breadcrumb, not a call-to-action
SOURCE_BOX_IN = (0.4, 5.2, 3.0, 0.3)  # bottom-left corner by default

# Table type (<table>): a native pptx table, themed *here* rather than by the Office
# theme — the renderer swaps in the built-in "No Style, No Grid" and writes every
# fill/border itself, so a row's contrast is a stated colour, not an inherited one.
# Row contrast comes two ways and they compose: `band` (zebra fill) and `rules`
# (horizontal lines). Defaults below are the house look: no header fill, a rule under
# the header, a faint zebra on alternate body rows.
TABLE_NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"   # built-in "No Style, No Grid"
TABLE_SZ_PT = 12
TABLE_ROW_H_IN = 0.42
TABLE_HEADER_FILL = "none"            # the header reads by weight + rule, not a slab
TABLE_HEADER_BOLD = True
TABLE_BAND_HEX = "F4F4F0"             # zebra fill on alternate body rows
TABLE_RULE_HEX = "CFCFC9"
TABLE_RULE_PT = 0.75
TABLE_HEADER_RULE_PT = 1.25           # the header rule reads heavier than a row rule
TABLE_PAD_IN = (0.10, 0.03, 0.10, 0.03)   # left, top, right, bottom cell margin
# ---------------------------------------------------------------------------

# --- the card family -------------------------------------------------------
# The card family — native <turn> cards (transcripts) and any image-side code/diff
# card renderer paired with them share these so they read as one object: same
# fill, border, radius, a small top-left context label, and a tight mono body.


@dataclass(frozen=True)
class CardStyle:
    """A per-actor card look: light fill, border colour, accent (label + border),
    and the accent border weight (pt). Frozen — a shared constant, not mutable
    state."""

    fill: str
    border: str
    accent: str
    border_pt: float


# Native + image shared geometry / type.
CARD_RADIUS_IN = 0.06
CARD_LABEL_SZ = 10                 # small top-left context label, pt
CARD_BODY_SZ = 12                  # mono body, pt
CARD_BODY_MONO = "Courier New"    # Slides-safe mono — Consolas isn't a Google Slides
                                  # font, so it silently substitutes a proportional
                                  # fallback there; image side maps to DejaVu Sans Mono
CARD_LABEL_HEX = "707070"         # muted label grey
CARD_LABEL_GAP_PT = 2             # space_after under the label paragraph
CARD_LINE_SPACING = 0.95          # tight body line spacing
CARD_STRIPE_PT = 2.25             # accent border weight on native turns
CARD_PAD_IN = (0.12, 0.08, 0.12, 0.08)   # left, top, right, bottom internal margin

# Actor -> CardStyle. `actor` is the STYLE key; the displayed label can differ (label=).
CARD_ACTORS: dict[str, CardStyle] = {
    "user":   CardStyle(fill="EFEFEC", border="CFCFC9", accent="555555", border_pt=1),
    "main":   CardStyle(fill="EFEFEC", border="CFCFC9", accent="555555", border_pt=1),
    "leaf":   CardStyle(fill="EAF0FB", border="9DB6E0", accent="2A52A0", border_pt=1),
    "agent":  CardStyle(fill="FBE0E0", border="A23A3A", accent="A23A3A", border_pt=1),
    "result": CardStyle(fill="FBE0E0", border="A23A3A", accent="A23A3A", border_pt=2),
}


def card_style(actor: str) -> CardStyle:
    """Resolve an actor name to its CardStyle, or raise with the allowed set.
    Raises ValueError (not LayoutError) to avoid an import cycle — the emitter
    translates it; layout_to_pptx imports look, never vice-versa."""
    try:
        return CARD_ACTORS[actor]
    except KeyError:
        raise ValueError(f"actor={actor!r} not in {sorted(CARD_ACTORS)}") from None


# Image-card palette, for an image-side code/diff card renderer that should match
# the native <turn> cards ('#'-prefixed hex, as matplotlib/PIL take it).
CODE_CARD_BG = "#FBFBF8"
CODE_CARD_BORDER = "#CFCFC9"
CODE_INK = "#24292E"
CODE_KW = "#9B2393"
CODE_STR = "#1A7F37"
CODE_COMMENT = "#8A8A8A"
CODE_WIRE = "#C0392B"
CODE_TAG = "#2A6FB0"
CODE_DIFF_ADD_FG = "#1A7F37"
CODE_DIFF_ADD_BG = "#E6F6EA"
CODE_DIFF_CTX = "#8A8A8A"
CODE_BLOCK_BG = "#EAF2FB"
# ---------------------------------------------------------------------------

# DrawingML / PresentationML namespaces, for the oxml fragments built by hand.
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def add_border(master: SlideMaster) -> None:
    """Append a no-fill framed rectangle to the master shape tree; on the master
    so the frame shows on every slide (incl. the title) from one definition."""
    off = BORDER_INSET
    sp = parse_xml(
        f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr>'
        '<p:cNvPr id="60" name="Slide Border"/><p:cNvSpPr/><p:nvPr/>'
        '</p:nvSpPr><p:spPr>'
        f'<a:xfrm><a:off x="{off}" y="{off}"/>'
        f'<a:ext cx="{SLIDE_W - 2 * off}" cy="{SLIDE_H - 2 * off}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/>'
        f'<a:ln w="{BORDER_WEIGHT}"><a:solidFill>'
        f'<a:srgbClr val="{BORDER_HEX}"/></a:solidFill></a:ln>'
        '</p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>'
    )
    sptree = master.element.find(qn("p:cSld")).find(qn("p:spTree"))
    sptree.append(sp)


def add_slide_number(slide: Slide) -> None:
    """Add a bottom-right slide number as an auto-updating field. The
    `<a:fld type="slidenum">` is what editors substitute with the live number;
    python-pptx has no field API, so the run is built as an oxml fragment."""
    box = slide.shapes.add_textbox(
        Emu(SLIDE_NUM_OFF[0]), Emu(SLIDE_NUM_OFF[1]),
        Emu(SLIDE_NUM_EXT[0]), Emu(SLIDE_NUM_EXT[1]),
    )
    para = box.text_frame.paragraphs[0]
    para.alignment = PP_ALIGN.RIGHT
    para._p.append(parse_xml(
        f'<a:fld xmlns:a="{_A}" id="{{B9D74E2A-1B0E-4D2C-9C1F-0A6F4E5C7A11}}" '
        'type="slidenum">'
        f'<a:rPr lang="en-US" sz="{SLIDE_NUM_SZ}"><a:solidFill>'
        f'<a:srgbClr val="{SLIDE_NUM_HEX}"/></a:solidFill></a:rPr>'
        '<a:t>2</a:t></a:fld>'
    ))


def outline_textbox(shape) -> None:  # noqa: ANN001 — pptx Shape, no public type alias
    """Add the faint editing-bounds outline to one content text box, unless it
    already carries a line (an author-set border, or a suppress_outline call,
    wins). Written as an oxml fragment so it doesn't depend on the wrapper type
    exposing `.line`. The renderer gives every shape explicit geometry, so —
    unlike the legacy path — there's nothing to materialise first."""
    if not TEXTBOX_OUTLINE_SHOW:             # outlines disabled deck-wide
        return
    sppr = shape._element.spPr
    if sppr.find(qn("a:ln")) is not None:   # author/explicit line already set
        return
    sppr.append(parse_xml(
        f'<a:ln xmlns:a="{_A}" w="{TEXTBOX_OUTLINE_WEIGHT}"><a:solidFill>'
        f'<a:srgbClr val="{TEXTBOX_OUTLINE_HEX}"/></a:solidFill></a:ln>'
    ))


def suppress_outline(shape) -> None:  # noqa: ANN001 — pptx Shape, no public type alias
    """Give a shape an explicit no-line so the faint chrome outline skips it
    (`outline="none"` in a layout file). An oxml fragment, so it's type-safe for
    any shape wrapper."""
    sppr = shape._element.spPr
    if sppr.find(qn("a:ln")) is None:
        sppr.append(parse_xml(f'<a:ln xmlns:a="{_A}"><a:noFill/></a:ln>'))


# In the a:tcPr schema the four edge lines come first, ahead of the cell fill; these
# are the elements a new a:lnL/R/T/B must be inserted before to stay schema-valid.
_TCPR_AFTER_LN = (
    "a:lnTlToBr", "a:lnBlToTr", "a:cell3D", "a:noFill", "a:solidFill", "a:gradFill",
    "a:blipFill", "a:pattFill", "a:grpFill", "a:headers", "a:extLst",
)
_CELL_EDGE = {"left": "a:lnL", "right": "a:lnR", "top": "a:lnT", "bottom": "a:lnB"}


def cell_border(cell, edge: str, hex_value: str, weight_pt: float) -> None:  # noqa: ANN001
    """Draw one edge of a table cell. python-pptx has no border API, so the
    `a:lnL/R/T/B` goes in as an oxml fragment — inserted ahead of the fill,
    because the tcPr schema fixes that order."""
    tag = _CELL_EDGE[edge]
    tcpr = cell._tc.get_or_add_tcPr()
    for existing in tcpr.findall(qn(tag)):
        tcpr.remove(existing)                # restate, don't stack
    tcpr.insert_element_before(parse_xml(
        f'<{tag} xmlns:a="{_A}" w="{int(Pt(weight_pt))}" cap="flat" cmpd="sng" '
        f'algn="ctr"><a:solidFill><a:srgbClr val="{hex_value.lstrip("#").upper()}"/>'
        f'</a:solidFill><a:prstDash val="solid"/></{tag}>'
    ), *_TCPR_AFTER_LN)


def unstyle_table(table) -> None:  # noqa: ANN001 — pptx Table, no public type alias
    """Strip the Office theme's table style (python-pptx defaults to a blue banded
    one) so the emitter's own fills and borders are the whole look. Turns off the
    style's first-row/banding parts and points the table at "No Style, No Grid"."""
    table.first_row = False
    table.horz_banding = False
    tbl_pr = table._tbl.find(qn("a:tblPr"))
    if tbl_pr is None:
        return
    style_id = tbl_pr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = parse_xml(f'<a:tableStyleId xmlns:a="{_A}"/>')
        tbl_pr.append(style_id)
    style_id.text = TABLE_NO_STYLE


def hex_to_rgb(value: str) -> RGBColor:
    """Parse a `RRGGBB` (or `#RRGGBB`) hex string to a python-pptx RGBColor."""
    return RGBColor.from_string(value.lstrip("#").upper())


def pt(size: float):  # noqa: ANN201 — pptx Length, no public type alias
    """Point size -> EMU Length (thin wrapper so callers don't import Pt)."""
    return Pt(size)
