#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-pptx>=1.0.2,<2", "typer>=0.12,<1"]
# ///
"""Fidelity check for the XML-layout -> pptx path: is the pptx a faithful
representation of the input, and will its text fit where it was placed?

Three signals:

1. **Element presence (the gate).** Re-count the layout, then assert every
   declared shape-producing element became a real shape on the right slide.
   Catches silent drops and surprise slide splits; deterministic, runs in CI.
2. **Text fit (advisory, `--strict` makes it a gate).** Estimate each text box's
   wrapped height from its runs' sizes and the paragraph line spacing, and warn
   when the text needs more room than the box has — the wrapped-title-into-the-
   subtitle failure a shape count can't see. A heuristic (average glyph widths),
   deliberately a little pessimistic. Also warns when a deck font isn't installed
   here, since the LibreOffice render then wraps earlier than PowerPoint/Slides.
3. **Visual render (the eyeball signal).** pptx -> PDF (LibreOffice headless) ->
   per-slide PNGs (poppler), so the author or an agent can confirm placement.
   Skipped with a note if `soffice`/`pdftoppm` are missing.

Exit code is non-zero iff element presence fails (or, with `--strict`, a text-fit
warning fired).
"""

from __future__ import annotations

import math
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

try:
    import typer
    from pptx import Presentation
    from pptx.util import Length
except ImportError as exc:  # a bare `python3` without the deps installed
    sys.exit(
        f"[fidelity] missing dependency: {exc.name}. Run via `uv run {Path(__file__).name}` "
        "(deps are declared inline) or `pip install python-pptx typer`."
    )

# The layout tags that must each yield exactly one shape in the output. <columns>
# / <column> are containers (they yield no shape of their own); their children
# are counted where they sit. Kept in sync with layout_to_pptx._EMITTERS. turn ->
# 1 shape (label + body paragraphs live inside it); table -> 1 graphic frame (its
# <tr>/<td> live inside it, so the walk never descends into one).
_SHAPE_TAGS = {
    "title", "text", "caption", "source", "bullets", "image", "shape", "turn", "table",
}

# Text-fit heuristic knobs. Average advance width per character as a fraction of
# the font size: Calibri/Carlito are narrow (~0.48em); the common Linux fallback
# (DejaVu Sans) is ~15% wider. Line height for spcPct 100% is ~1.2x the size.
_CHAR_W = {"regular": 0.50, "bold": 0.53, "mono": 0.60}
_LINE_H = 1.2
_SLACK = 1.05          # tolerate 5% over before warning
_EMU_PER_PT = 12700

# Metric-compatible substitutes: a render with one of these matches PowerPoint's
# line breaks; anything else wraps differently.
_METRIC_OK = {
    "calibri": ({"calibri", "carlito"}, "install Carlito (Debian/Ubuntu: fonts-crosextra-carlito)"),
    "courier new": ({"couriernew", "liberationmono", "cousine"}, "install Liberation Mono"),
}


def _count_declared(layout_path: Path) -> tuple[int, list[tuple[int, str]]]:
    """Walk the layout file and count declared shape-producing elements,
    returning (total, per-slide list of (slide_idx, tag)). Mirrors the
    renderer's traversal (descending into <columns>/<column>)."""
    root = ET.parse(layout_path).getroot()
    per_slide: list[tuple[int, str]] = []

    def walk(el: ET.Element, idx: int) -> None:
        for child in el:
            if child.tag in ("columns", "column"):
                walk(child, idx)          # descend; containers emit nothing
            elif child.tag in _SHAPE_TAGS:
                per_slide.append((idx, child.tag))

    for idx, sl in enumerate(root.findall("slide")):
        walk(sl, idx)
    return len(per_slide), per_slide


def _is_slide_number(shape) -> bool:  # noqa: ANN001 — pptx Shape
    """True for the auto-added slide-number box (carries a slidenum field)."""
    if not shape.has_text_frame:
        return False
    return "slidenum" in shape.text_frame._txBody.xml


def _shapes_per_slide(prs) -> list[int]:  # noqa: ANN001 — pptx Presentation
    """Author shapes on each slide, excluding the renderer's chrome (the slide
    number; the master border isn't a slide shape)."""
    counts: list[int] = []
    for slide in prs.slides:
        counts.append(sum(1 for s in slide.shapes if not _is_slide_number(s)))
    return counts


def element_presence(layout_path: Path, prs) -> tuple[bool, str]:  # noqa: ANN001
    """Assert every declared element maps to a shape: per-slide declared count
    must equal per-slide author-shape count in the output. Returns (ok, report)."""
    total, per_slide = _count_declared(layout_path)
    declared_by_slide: dict[int, int] = {}
    for idx, _tag in per_slide:
        declared_by_slide[idx] = declared_by_slide.get(idx, 0) + 1
    actual = _shapes_per_slide(prs)
    lines: list[str] = []
    ok = True
    n_slides = max(len(actual), (max(declared_by_slide) + 1) if declared_by_slide else 0)
    for idx in range(n_slides):
        want = declared_by_slide.get(idx, 0)
        got = actual[idx] if idx < len(actual) else 0
        mark = "ok" if want == got else "MISMATCH"
        if want != got:
            ok = False
        lines.append(f"  slide {idx + 1}: declared={want} shapes={got} [{mark}]")
    head = (
        f"element-presence: {total} declared shape-elements, "
        f"{sum(actual)} author shapes in pptx — {'PASS' if ok else 'FAIL'}"
    )
    return ok, head + "\n" + "\n".join(lines)


def _text_shapes(prs):  # noqa: ANN001, ANN201 — yields (slide_no, shape)
    for idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame and not _is_slide_number(shape):
                yield idx, shape


def _para_need_pt(para) -> float:  # noqa: ANN001 — pptx _Paragraph
    """Estimated height of one paragraph once wrapped: (lines x line height) +
    space-after, given the frame width passed via `para._fit_width_pt`."""
    width_pt = 0.0
    size_max = 0.0
    for run in para.runs:
        size = run.font.size.pt if run.font.size is not None else 14.0
        name = (run.font.name or "").lower()
        is_mono = any(m in name for m in ("courier", "consolas", "mono"))
        key = "mono" if is_mono else ("bold" if run.font.bold else "regular")
        width_pt += len(run.text) * size * _CHAR_W[key]
        size_max = max(size_max, size)
    size_max = size_max or 14.0
    spacing = para.line_spacing
    if isinstance(spacing, Length):          # exact points (spcPts); Length is an int subclass
        line_h = spacing.pt
    elif spacing is not None:                # a multiple (spcPct)
        line_h = size_max * _LINE_H * float(spacing)
    else:
        line_h = size_max * _LINE_H
    lines = max(1, math.ceil(width_pt / para._fit_width_pt)) if width_pt else 1
    after = para.space_after.pt if para.space_after is not None else 0.0
    return lines * line_h + after


def text_fit(prs) -> list[str]:  # noqa: ANN001 — pptx Presentation
    """Warn for every text shape whose estimated wrapped height exceeds its box."""
    warnings: list[str] = []
    for idx, shape in _text_shapes(prs):
        tf = shape.text_frame
        if not tf.text.strip():
            continue
        fit_w = shape.width / _EMU_PER_PT - tf.margin_left.pt - tf.margin_right.pt
        if fit_w <= 0:
            continue
        need = tf.margin_top.pt + tf.margin_bottom.pt
        for para in tf.paragraphs:
            para._fit_width_pt = fit_w
            need += _para_need_pt(para)
        have = shape.height / _EMU_PER_PT
        if need > have * _SLACK + 4:
            snippet = tf.text.strip().replace("\n", " ")[:48]
            warnings.append(
                f"  slide {idx}: '{snippet}' needs ~{need / 72:.2f}in but its box is "
                f"{have / 72:.2f}in high — text will spill past the box"
            )
    return warnings


def font_substitutions(prs) -> list[str]:  # noqa: ANN001 — pptx Presentation
    """Warn when a font the deck uses isn't installed here (fontconfig picks a
    non-metric-compatible substitute, so the render wraps differently)."""
    fc_match = shutil.which("fc-match")
    if fc_match is None:
        return []
    fonts: set[str] = set()
    for _idx, shape in _text_shapes(prs):
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if run.font.name:
                    fonts.add(run.font.name)
    notes: list[str] = []
    for font in sorted(fonts):
        out = subprocess.run([fc_match, font], capture_output=True, text=True).stdout
        matched = out.split('"')[1] if '"' in out else out.strip()
        ok, hint = _METRIC_OK.get(font.lower(), ({font.lower().replace(" ", "")}, "install it"))
        if matched.lower().replace(" ", "") not in ok:
            notes.append(
                f"  font '{font}' renders as '{matched}' on this machine — {hint} so the "
                "PNG/PDF render breaks lines where PowerPoint and Google Slides do"
            )
    return notes


def _soffice() -> str | None:
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found is None and Path("/Applications/LibreOffice.app/Contents/MacOS/soffice").exists():
        found = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return found


def render_pngs(pptx_path: Path, out_dir: Path, *, dpi: int = 110) -> list[Path]:
    """Render the pptx to per-slide PNGs (soffice -> pdf -> pdftoppm). Returns
    the PNG paths, or [] if the binaries are missing (caller notes the skip)."""
    soffice = _soffice()
    pdftoppm = shutil.which("pdftoppm")
    if not (soffice and pdftoppm):
        return []
    out_dir.mkdir(parents=True, exist_ok=True)
    # A private profile dir avoids clashing with an interactive LibreOffice.
    profile = out_dir / "_lo_profile"
    subprocess.run(
        [soffice, "--headless", f"-env:UserInstallation=file://{profile}",
         "--convert-to", "pdf", "--outdir", str(out_dir), str(pptx_path)],
        check=True, capture_output=True, timeout=180,
    )
    pdf = out_dir / (pptx_path.stem + ".pdf")
    if not pdf.is_file():
        return []
    subprocess.run(
        [pdftoppm, "-png", "-r", str(dpi), str(pdf), str(out_dir / "slide")],
        check=True, capture_output=True, timeout=120,
    )
    return sorted(out_dir.glob("slide-*.png"))


def main(
    layout: Path = typer.Argument(..., help="the .xml layout file"),
    pptx: Path = typer.Option(None, "--pptx", help="the produced .pptx (default: <layout>.pptx)"),
    png_dir: Path = typer.Option(None, "--png-dir", help="write per-slide PNGs here for a visual check"),
    dpi: int = typer.Option(110, help="PNG render resolution"),
    strict: bool = typer.Option(False, "--strict", help="fail on text-fit warnings too"),
) -> None:
    """Check a produced pptx is faithful to its XML layout (presence, text fit, render)."""
    pptx_path = pptx or layout.with_suffix(".pptx")
    if not layout.is_file():
        typer.echo(f"[fidelity] layout not found: {layout}", err=True)
        raise typer.Exit(2)
    if not pptx_path.is_file():
        typer.echo(f"[fidelity] pptx not found: {pptx_path} (render it first)", err=True)
        raise typer.Exit(2)
    prs = Presentation(str(pptx_path))
    ok, report = element_presence(layout, prs)
    typer.echo(report)
    fit_warnings = text_fit(prs)
    for line in fit_warnings + font_substitutions(prs):
        typer.echo(f"[fidelity] warn:\n{line}")
    if not fit_warnings:
        typer.echo("text-fit: every text box has room for its text (heuristic)")
    if png_dir is not None:
        pngs = render_pngs(pptx_path, png_dir, dpi=dpi)
        if pngs:
            typer.echo(f"[fidelity] rendered {len(pngs)} PNGs to {png_dir}")
        else:
            typer.echo("[fidelity] visual render skipped — soffice/pdftoppm missing (element-presence is the gate)")
    failed = not ok or (strict and fit_warnings)
    raise typer.Exit(1 if failed else 0)


if __name__ == "__main__":
    typer.run(main)
